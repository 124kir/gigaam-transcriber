"""Парсинг таймкодов и экспорт транскрипта в SRT/VTT/JSON/DOCX."""
import json
from typing import List, Tuple


def format_srt_time(seconds: float) -> str:
    """Форматирование времени для SRT (часы:минуты:секунды,миллисекунды)"""
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    millis = int((seconds - int(seconds)) * 1000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d},{millis:03d}"

def format_vtt_time(seconds: float) -> str:
    """Форматирование времени для WebVTT (часы:минуты:секунды.миллисекунды)"""
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    millis = int((seconds - int(seconds)) * 1000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}.{millis:03d}"

def parse_transcript_timestamps(text: str) -> List[Tuple[float, float, str]]:
    """Парсит транскрипт с таймкодами [start-end] / [start-ends] / SPEAKER_XX [...]."""
    import re
    segments = []
    # Форматы, которые генерирует приложение:
    #   run_gigaam_simple:      "[12.3-45.6s] текст"
    #   run_gigaam_with_diar:   "SPEAKER_01 [12.3-45.6s] текст"
    # Суффикс 's' и префикс SPEAKER_XX — опциональны.
    pattern = r'(?:(SPEAKER_\S+)\s+)?\[(\d+\.?\d*)-(\d+\.?\d*)s?\]\s*(.*)'
    for match in re.finditer(pattern, text):
        speaker, start, end, content = match.group(1), match.group(2), match.group(3), match.group(4).strip()
        if speaker:
            content = f"{speaker} {content}".strip()
        segments.append((float(start), float(end), content))
    return segments

def export_srt(transcript: str, output_path: str) -> bool:
    """Экспорт транскрипта в формат SRT"""
    try:
        segments = parse_transcript_timestamps(transcript)
        if not segments:
            # Если нет таймкодов, создаём из строк
            lines = [l.strip() for l in transcript.splitlines() if l.strip()]
            segments = [(i*3, i*3+3, line) for i, line in enumerate(lines)]
        
        with open(output_path, 'w', encoding='utf-8') as f:
            for i, (start, end, text) in enumerate(segments, 1):
                f.write(f"{i}\n")
                f.write(f"{format_srt_time(start)} --> {format_srt_time(end)}\n")
                f.write(f"{text}\n\n")
        return True
    except Exception as e:
        print(f"Ошибка экспорта SRT: {e}")
        return False

def export_vtt(transcript: str, output_path: str) -> bool:
    """Экспорт транскрипта в формат WebVTT"""
    try:
        segments = parse_transcript_timestamps(transcript)
        if not segments:
            lines = [l.strip() for l in transcript.splitlines() if l.strip()]
            segments = [(i*3, i*3+3, line) for i, line in enumerate(lines)]
        
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write("WEBVTT\n\n")
            for start, end, text in segments:
                f.write(f"{format_vtt_time(start)} --> {format_vtt_time(end)}\n")
                f.write(f"{text}\n\n")
        return True
    except Exception as e:
        print(f"Ошибка экспорта VTT: {e}")
        return False

def export_json(transcript: str, output_path: str) -> bool:
    """Экспорт транскрипта в формат JSON"""
    try:
        import re
        segments = parse_transcript_timestamps(transcript)
        if not segments:
            lines = [l.strip() for l in transcript.splitlines() if l.strip()]
            segments = [(i*3, i*3+3, line) for i, line in enumerate(lines)]
        
        data = {
            "transcript": [
                {"start": start, "end": end, "text": text}
                for start, end, text in segments
            ],
            "metadata": {
                "total_segments": len(segments),
                "duration": segments[-1][1] if segments else 0
            }
        }
        
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        return True
    except Exception as e:
        print(f"Ошибка экспорта JSON: {e}")
        return False

def export_docx(transcript: str, output_path: str, summary: str = "") -> bool:
    """Экспорт транскрипта в формат DOCX"""
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        print("python-docx не установлен. Установите: pip install python-docx")
        return False
    
    try:
        doc = Document()
        
        # Заголовок
        title = doc.add_heading('Транскрипт аудио', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Транскрипт
        doc.add_heading('Текст транскрипта', level=1)
        
        segments = parse_transcript_timestamps(transcript)
        if not segments:
            # Простой текст без таймкодов
            for line in transcript.splitlines():
                if line.strip():
                    doc.add_paragraph(line)
        else:
            # С таймкодами
            for start, end, text in segments:
                time_str = f"[{format_srt_time(start)} --> {format_srt_time(end)}]"
                p = doc.add_paragraph()
                p.add_run(time_str).bold = True
                p.add_run(f" {text}")
        
        # Суммаризация, если есть
        if summary:
            doc.add_page_break()
            doc.add_heading('Краткая сводка', level=1)
            doc.add_paragraph(summary)
        
        doc.save(output_path)
        return True
    except Exception as e:
        print(f"Ошибка экспорта DOCX: {e}")
        return False
